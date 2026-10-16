from docx import Document
from pathlib import Path
from io import BytesIO

def replace_text(document, old_text, new_text):
    for paragraph in document.paragraphs:
        if old_text in paragraph.text:
            paragraph.text = paragraph.text.replace(old_text, new_text)

def export_to_word(text, file_path, template_path, replacements):
    document = Document(template_path)
    for old_text, new_text in replacements.items():
        replace_text(document, old_text, new_text)
    document.save(file_path)

def create_word_document(template_path, replacements):
    document = Document(template_path)
    document = Document(template_path)
    for old_text, new_text in replacements.items():
        replace_text(document, old_text, new_text)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return buffer

    